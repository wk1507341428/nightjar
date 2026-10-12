#!/usr/bin/env python3
"""对比小程序聚合品牌货源与得物 SKU 价格，并生成盈利分析 Excel。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import random
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


CATALOG_API = "https://aiopro.fvo2o.com/api/h5app/wxapp"
CATALOG_APP_ID = "wx8fa740aa841dafb4"
SIDEJOB_API = "http://127.0.0.1:9300"
DEWU_API = "https://stark.dewu.com/api/v1/h5"
DEWU_SIGN_SALT = "048a9c4943398714b356a696503d2d36"
TARGET_PROFILES = ("ADIDAS-合并", "ANTA KIDS - 聚合")
REGION_NAMES = {"2": "京津", "3": "上海", "4": "广佛", "5": "成都", "6": "武汉", "7": "重庆"}


class ValidationRequired(RuntimeError):
    pass


def request_json(url: str, *, headers=None, body=None, timeout=30):
    data = None if body is None else json.dumps(body, ensure_ascii=False, separators=(",", ":")).encode()
    request = urllib.request.Request(url, data=data, headers=headers or {}, method="POST" if body is not None else "GET")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.status, json.load(response)
    except urllib.error.HTTPError as error:
        try:
            payload = json.load(error)
        except Exception:
            payload = {}
        if error.code == 485:
            raise ValidationRequired("HTTP 485：得物触发校验，请提供新的完整 curl 后继续") from error
        return error.code, payload


def curl_json(url: str, *, headers: dict, body: dict, timeout=30):
    command = ["curl", "--silent", "--show-error", "--url", url, "--max-time", str(timeout)]
    for name, value in headers.items():
        if name == "Cookie":
            command.extend(["-b", value])
        else:
            command.extend(["-H", f"{name}: {value}"])
    command.extend(["--data-raw", json.dumps(body, ensure_ascii=False, separators=(",", ":")), "--write-out", "\n__HTTP_STATUS__:%{http_code}"])
    completed = subprocess.run(command, capture_output=True, text=True, timeout=timeout + 5, check=False)
    if completed.returncode:
        raise RuntimeError(completed.stderr.strip() or f"curl 退出码 {completed.returncode}")
    response_text, marker, status_text = completed.stdout.rpartition("\n__HTTP_STATUS__:")
    if not marker:
        raise RuntimeError("curl 响应缺少状态码")
    try:
        payload = json.loads(response_text)
    except json.JSONDecodeError as error:
        raise RuntimeError("curl 响应不是 JSON") from error
    return int(status_text.strip()), payload


def compact_value(value):
    if value is None or value == "":
        return ""
    if isinstance(value, list):
        return ",".join(compact_value(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def sign_body(body: dict) -> dict:
    source = "".join(f"{key}{compact_value(body[key])}" for key in sorted(body)) + DEWU_SIGN_SALT
    return {**body, "sign": hashlib.md5(source.encode()).hexdigest()}


class DewuClient:
    def __init__(self, min_interval: float, max_interval: float):
        self.min_interval = max(1.0, min_interval)
        self.max_interval = max(self.min_interval, max_interval)
        self.last_request_at = 0.0
        self.current_keyword = ""
        token = os.environ.get("DEWU_PASSPORT_TOKEN", "")
        security_key = os.environ.get("DEWU_SK", "")
        if not token or not security_key:
            raise RuntimeError("请设置 DEWU_PASSPORT_TOKEN 和 DEWU_SK")
        sys_code = os.environ.get("DEWU_SYS_CODE", "DEWU_MERCHANT_PLATFORM_DU_USER_T")
        cookie = os.environ.get("DEWU_COOKIE") or f"sk={security_key}; clientId=stark; sysCode={sys_code}; channel=pc; mchToken={token}"
        self.headers = {
            "Accept": "application/json",
            "Accept-Language": "zh-CN,zh;q=0.9",
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "Content-Type": "application/json;charset=UTF-8",
            "Cookie": cookie,
            "Origin": "https://stark.dewu.com",
            "Pragma": "no-cache",
            "Priority": "u=1, i",
            "Referer": "https://stark.dewu.com/main/businessApp/opportunitySearch?keyword=LC1295&from=%2Fmain%2FbusinessApp%2FmarketRanking",
            "SK": security_key,
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-origin",
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/153.0.0.0 Safari/537.36",
            "appid": "h5",
            "channel": "pc",
            "clientId": "stark",
            "passportToken": token,
            "platform": "h5",
            "sec-ch-ua": '"Google Chrome";v="153", "Not_A Brand";v="8", "Chromium";v="153"',
            "sec-ch-ua-mobile": "?0",
            "sec-ch-ua-platform": '"macOS"',
            "sysCode": sys_code,
        }

    def post(self, route: str, body: dict):
        interval = random.uniform(self.min_interval, self.max_interval)
        delay = interval - (time.monotonic() - self.last_request_at)
        if delay > 0:
            time.sleep(delay)
        self.last_request_at = time.monotonic()
        keyword = str(body.get("keyword") or self.current_keyword or "").strip()
        headers = dict(self.headers)
        headers["traceparent"] = os.environ.get("DEWU_TRACEPARENT") or f"00-{os.urandom(16).hex()}-{os.urandom(8).hex()}-01"
        if keyword:
            headers["Referer"] = (
                "https://stark.dewu.com/main/businessApp/opportunitySearch?"
                f"keyword={urllib.parse.quote(keyword)}&from=%2Fmain%2FbusinessApp%2FmarketRanking"
            )
        status, payload = curl_json(f"{DEWU_API}{route}", headers=headers, body=sign_body(body))
        message = str(payload.get("msg") or payload.get("message") or "")
        if status == 485 or payload.get("code") == 485 or any(word in message for word in ("验证码", "人机校验", "触发校验")):
            raise ValidationRequired(f"得物需要校验：{message or status}")
        if status != 200 or payload.get("code") != 200:
            raise RuntimeError(f"{route} 返回 {status}/{payload.get('code')}：{message}")
        return payload.get("data") or {}

    def query_item(self, item_no: str):
        self.current_keyword = item_no
        search = self.post("/merchant-tools/productSelect/search/getSearchProductPage", {
            "keyword": item_no,
            "biddingType": "0",
            "sortCode": "comprehensive",
            "pageCode": "SEARCH",
            "pageNum": 1,
            "pageSize": 14,
        })
        exact = next((item for item in search.get("contents", []) if normalize_item_no(item.get("articleNumber")) == item_no), None)
        if not exact:
            return {"status": "product_not_found"}
        detail = self.post("/biz/bid/merchant/aggregate/query/bidConfirmToBp", {
            "spuId": exact.get("spuId"),
            "bizSceneCode": 110,
            "bidTypes": ["0", "1", "15", "5"],
        })
        sku_rows = []
        for sku in detail.get("commodityInfo", {}).get("skuPropInfos", []):
            props = sku.get("props") or []
            size = next((str(prop.get("value") or "") for prop in props if prop.get("name") == "尺码"), str(sku.get("fullProp") or ""))
            sku_rows.append({
                "dewu_sku_id": str(sku.get("skuId") or ""),
                "dewu_size": size.strip(),
                "platform_min_price_cents": number(sku.get("platformMinPrice")),
                "estimated_net_cents": number(sku.get("estimateNetPrice")),
                "monthly_sales": number((sku.get("payOrderCntTotalL30d") or {}).get("value")),
                "bidding_auth": sku.get("biddingAuth") is not False,
                "dewu_tags": [tag.get("tagName") for tag in sku.get("skuTags", []) if tag.get("tagName")],
            })
        return {
            "status": "ok",
            "spu_id": str(exact.get("spuId") or ""),
            "dewu_title": exact.get("productName") or "",
            "total_monthly_sales": number(exact.get("platSalesCnt30d")),
            "skus": sku_rows,
        }


class BrowserDewuClient(DewuClient):
    """在已登录的真实浏览器页面中调用得物接口。"""

    def __init__(self, min_interval: float, max_interval: float, session: str, verify_timeout: int):
        self.min_interval = max(1.0, min_interval)
        self.max_interval = max(self.min_interval, max_interval)
        self.last_request_at = 0.0
        self.current_keyword = ""
        self.session = session
        self.verify_timeout = max(30, verify_timeout)

    def browser_eval(self, expression: str):
        completed = subprocess.run(
            ["agent-browser", "--session", self.session, "eval", expression],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        if completed.returncode:
            raise RuntimeError(completed.stderr.strip() or completed.stdout.strip() or "浏览器脚本执行失败")
        output = completed.stdout.strip()
        try:
            value = json.loads(output)
            return json.loads(value) if isinstance(value, str) else value
        except json.JSONDecodeError as error:
            raise RuntimeError(f"无法解析浏览器返回值：{output[:200]}") from error

    def wait_for_verification(self, session_id: str):
        print("得物触发校验，请在已打开的浏览器窗口完成拼图；完成后会自动继续。", flush=True)
        deadline = time.monotonic() + self.verify_timeout
        while time.monotonic() < deadline:
            state = self.browser_eval(
                "JSON.stringify({verified:window.__sidejobDewuVerified===true,"
                "sessionId:window.__sidejobDewuVerifySession||''})"
            )
            if state.get("verified") and state.get("sessionId") == session_id:
                print("校验已完成，重试当前货号。", flush=True)
                return
            time.sleep(2)
        raise ValidationRequired(f"等待浏览器校验超时（sessionId={session_id}）")

    def post(self, route: str, body: dict):
        interval = random.uniform(self.min_interval, self.max_interval)
        delay = interval - (time.monotonic() - self.last_request_at)
        if delay > 0:
            time.sleep(delay)
        self.last_request_at = time.monotonic()
        signed_body = sign_body(body)
        keyword = str(body.get("keyword") or self.current_keyword or "").strip()
        page_path = (
            "/main/businessApp/opportunitySearch?"
            f"keyword={urllib.parse.quote(keyword)}&from=%2Fmain%2FbusinessApp%2FmarketRanking"
        )
        expression = f"""
(async()=>{{
  const body={json.dumps(signed_body, ensure_ascii=False, separators=(',', ':'))};
  const route={json.dumps(route)};
  const pagePath={json.dumps(page_path)};
  const cookies=Object.fromEntries(document.cookie.split('; ').filter(Boolean).map(value=>{{
    const index=value.indexOf('=');
    return [decodeURIComponent(value.slice(0,index)),decodeURIComponent(value.slice(index+1))];
  }}));
  const token=cookies.mchToken||'';
  const securityKey=window.dw&&window.dw.getSk?String(window.dw.getSk()||''):'';
  if(!token||!securityKey) return JSON.stringify({{bridgeError:'浏览器登录态或 SK 缺失'}});
  history.replaceState(null,'',pagePath);
  const response=await fetch('/api/v1/h5'+route,{{
    method:'POST',
    credentials:'include',
    headers:{{
      Accept:'application/json',
      'Content-Type':'application/json;charset=UTF-8',
      appid:'h5',
      channel:'pc',
      clientId:cookies.clientId||'stark',
      passportToken:token,
      platform:'h5',
      SK:securityKey,
      sysCode:cookies.sysCode||'DEWU_MERCHANT_PLATFORM_DU_USER_T'
    }},
    body:JSON.stringify(body)
  }});
  let payload={{}};
  try {{ payload=await response.json(); }} catch(error) {{ payload={{msg:String(error)}}; }}
  if(response.status===485&&payload&&payload.data&&payload.data.type===2){{
    const sessionId=String(payload.data.sessionId||'');
    window.__sidejobDewuVerified=false;
    window.__sidejobDewuVerifySession=sessionId;
    if(window.dewuCap) window.dewuCap({{
      sessionId,
      sumei:'',
      lang:'',
      success:()=>{{window.__sidejobDewuVerified=true;}}
    }});
  }}
  return JSON.stringify({{status:response.status,payload}});
}})()
""".strip()
        result = self.browser_eval(expression)
        if result.get("bridgeError"):
            raise RuntimeError(result["bridgeError"])
        status = int(result.get("status") or 0)
        payload = result.get("payload") or {}
        message = str(payload.get("msg") or payload.get("message") or "")
        if status == 485 or payload.get("code") == 485:
            session_id = str((payload.get("data") or {}).get("sessionId") or "")
            if not session_id:
                raise ValidationRequired(f"得物需要校验：{message or status}")
            self.wait_for_verification(session_id)
            return self.post(route, body)
        if status != 200 or payload.get("code") != 200:
            raise RuntimeError(f"{route} 返回 {status}/{payload.get('code')}：{message}")
        return payload.get("data") or {}

    def run_batch(self, requests: list[dict]) -> list[dict]:
        """单次进入浏览器，按间隔顺序执行一批同源请求。"""
        pending = list(requests)
        completed = []
        while pending:
            browser_requests = [
                {
                    "item_no": request["item_no"],
                    "route": request["route"],
                    "body": sign_body(request["body"]),
                }
                for request in pending
            ]
            interval_ms = int(random.uniform(self.min_interval, self.max_interval) * 1000)
            expression = f"""
(async()=>{{
  const requests={json.dumps(browser_requests, ensure_ascii=False, separators=(',', ':'))};
  const intervalMs={interval_ms};
  const cookies=Object.fromEntries(document.cookie.split('; ').filter(Boolean).map(value=>{{
    const index=value.indexOf('=');
    return [decodeURIComponent(value.slice(0,index)),decodeURIComponent(value.slice(index+1))];
  }}));
  const token=cookies.mchToken||'';
  const securityKey=window.dw&&window.dw.getSk?String(window.dw.getSk()||''):'';
  if(!token||!securityKey) return JSON.stringify({{bridgeError:'浏览器登录态或 SK 缺失'}});
  const results=[];
  const sleep=ms=>new Promise(resolve=>setTimeout(resolve,ms));
  for(let index=0;index<requests.length;index++){{
    const request=requests[index];
    if(index>0) await sleep(intervalMs);
    history.replaceState(null,'','/main/businessApp/opportunitySearch?keyword='+encodeURIComponent(request.item_no)+'&from=%2Fmain%2FbusinessApp%2FmarketRanking');
    const response=await fetch('/api/v1/h5'+request.route,{{
      method:'POST',
      credentials:'include',
      headers:{{
        Accept:'application/json',
        'Content-Type':'application/json;charset=UTF-8',
        appid:'h5',channel:'pc',clientId:cookies.clientId||'stark',
        passportToken:token,platform:'h5',SK:securityKey,
        sysCode:cookies.sysCode||'DEWU_MERCHANT_PLATFORM_DU_USER_T'
      }},
      body:JSON.stringify(request.body)
    }});
    let payload={{}};
    try {{ payload=await response.json(); }} catch(error) {{ payload={{msg:String(error)}}; }}
    results.push({{item_no:request.item_no,status:response.status,payload}});
    if(response.status===485&&payload&&payload.data&&payload.data.type===2){{
      const sessionId=String(payload.data.sessionId||'');
      window.__sidejobDewuVerified=false;
      window.__sidejobDewuVerifySession=sessionId;
      if(window.dewuCap) window.dewuCap({{
        sessionId,sumei:'',lang:'',
        success:()=>{{window.__sidejobDewuVerified=true;}}
      }});
      break;
    }}
  }}
  return JSON.stringify({{results}});
}})()
""".strip()
            result = self.browser_eval(expression)
            if result.get("bridgeError"):
                raise RuntimeError(result["bridgeError"])
            batch_results = result.get("results") or []
            challenge = next(
                (
                    row
                    for row in batch_results
                    if row.get("status") == 485 or (row.get("payload") or {}).get("code") == 485
                ),
                None,
            )
            successful_count = len(batch_results) - (1 if challenge else 0)
            completed.extend(batch_results[:successful_count])
            pending = pending[successful_count:]
            if challenge:
                session_id = str(((challenge.get("payload") or {}).get("data") or {}).get("sessionId") or "")
                if not session_id:
                    raise ValidationRequired("得物需要校验，但响应未返回 sessionId")
                self.wait_for_verification(session_id)
                continue
            if successful_count == 0 and pending:
                raise RuntimeError("浏览器批次未返回任何结果")
        return completed

    def query_items(self, item_nos: list[str]) -> dict[str, dict]:
        search_rows = self.run_batch([
            {
                "item_no": item_no,
                "route": "/merchant-tools/productSelect/search/getSearchProductPage",
                "body": {
                    "keyword": item_no,
                    "biddingType": "0",
                    "sortCode": "comprehensive",
                    "pageCode": "SEARCH",
                    "pageNum": 1,
                    "pageSize": 14,
                },
            }
            for item_no in item_nos
        ])
        results = {}
        exact_products = {}
        for row in search_rows:
            item_no = row["item_no"]
            payload = row.get("payload") or {}
            if row.get("status") != 200 or payload.get("code") != 200:
                results[item_no] = {"status": "error", "error": str(payload.get("msg") or row.get("status"))}
                continue
            search = payload.get("data") or {}
            exact = next(
                (item for item in search.get("contents", []) if normalize_item_no(item.get("articleNumber")) == item_no),
                None,
            )
            if not exact:
                results[item_no] = {"status": "product_not_found"}
            else:
                exact_products[item_no] = exact

        detail_rows = self.run_batch([
            {
                "item_no": item_no,
                "route": "/biz/bid/merchant/aggregate/query/bidConfirmToBp",
                "body": {
                    "spuId": exact.get("spuId"),
                    "bizSceneCode": 110,
                    "bidTypes": ["0", "1", "15", "5"],
                },
            }
            for item_no, exact in exact_products.items()
        ]) if exact_products else []
        for row in detail_rows:
            item_no = row["item_no"]
            payload = row.get("payload") or {}
            exact = exact_products[item_no]
            if row.get("status") != 200 or payload.get("code") != 200:
                results[item_no] = {"status": "error", "error": str(payload.get("msg") or row.get("status"))}
                continue
            detail = payload.get("data") or {}
            sku_rows = []
            for sku in detail.get("commodityInfo", {}).get("skuPropInfos", []):
                props = sku.get("props") or []
                size = next(
                    (str(prop.get("value") or "") for prop in props if prop.get("name") == "尺码"),
                    str(sku.get("fullProp") or ""),
                )
                sku_rows.append({
                    "dewu_sku_id": str(sku.get("skuId") or ""),
                    "dewu_size": size.strip(),
                    "platform_min_price_cents": number(sku.get("platformMinPrice")),
                    "estimated_net_cents": number(sku.get("estimateNetPrice")),
                    "monthly_sales": number((sku.get("payOrderCntTotalL30d") or {}).get("value")),
                    "bidding_auth": sku.get("biddingAuth") is not False,
                    "dewu_tags": [tag.get("tagName") for tag in sku.get("skuTags", []) if tag.get("tagName")],
                })
            results[item_no] = {
                "status": "ok",
                "spu_id": str(exact.get("spuId") or ""),
                "dewu_title": exact.get("productName") or "",
                "total_monthly_sales": number(exact.get("platSalesCnt30d")),
                "skus": sku_rows,
            }
        return results


def number(value):
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def normalize_item_no(value):
    return str(value or "").strip().upper()


def normalize_spec(value):
    return str(value or "").strip().upper().replace(" ", "").replace("⅓", "1/3").replace("⅔", "2/3")


def first_positive(*values):
    for value in values:
        result = number(value)
        if result > 0:
            return result
    return 0


def load_profiles():
    _, payload = request_json(f"{SIDEJOB_API}/api/catalog/brand-maintenance", headers={"Accept": "application/json"})
    profiles = {profile.get("name"): profile for profile in payload.get("profiles", [])}
    missing = [name for name in TARGET_PROFILES if name not in profiles]
    if missing:
        raise RuntimeError(f"找不到聚合品牌：{', '.join(missing)}")
    return [profiles[name] for name in TARGET_PROFILES]


def catalog_page(member, page):
    query = urllib.parse.urlencode({
        "regionauth_id": member["regionId"], "distributor_id": member["distributorId"],
        "page": page, "pageSize": 50, "keywords": "", "approve_status": "onsale,only_show",
        "item_type": "normal", "is_point": "false", "company_id": 1,
    })
    _, payload = request_json(f"{CATALOG_API}/goods/items?{query}", headers={"Accept": "application/json", "authorizer-appid": CATALOG_APP_ID})
    data = payload.get("data") or {}
    if number(data.get("status_code")) >= 400:
        raise RuntimeError(data.get("message") or "小程序商品列表读取失败")
    return data.get("list") or [], number(data.get("total_count"))


def load_profile_products(profile):
    products = []
    member_counts = []
    for member in profile.get("members", []):
        first, total = catalog_page(member, 1)
        current = list(first)
        for page in range(2, (total + 49) // 50 + 1):
            current.extend(catalog_page(member, page)[0])
        region_name = REGION_NAMES.get(str(member["regionId"]), str(member["regionId"]))
        for product in current:
            product["_region_id"] = str(member["regionId"])
            product["_region_name"] = region_name
            product["_source_name"] = member.get("sourceName") or ""
        products.extend(current)
        member_counts.append({**member, "regionName": region_name, "productCount": len(current)})
        print(f"  {region_name} {member.get('sourceName')}: {len(current)} 个商品", flush=True)
    return products, member_counts


def source_skus(brand, products):
    rows = []
    for product in products:
        item_no = normalize_item_no(product.get("item_no"))
        if not item_no:
            continue
        product_price = first_positive(product.get("activity_price"), product.get("act_price"), product.get("price"))
        category = "/".join(entry.get("category_name") for entry in product.get("item_category_main", []) if entry.get("category_name"))
        for sku in product.get("spec_items") or []:
            stock = number(sku.get("store"))
            if stock <= 0 or str(sku.get("approve_status") or "").lower() not in ("onsale", "instock"):
                continue
            specs = sku.get("item_spec") or []
            preferred = next((entry for entry in specs if any(word in str(entry.get("spec_name") or "") for word in ("尺码", "尺寸", "型号", "规格"))), None)
            fallback = next((entry for entry in specs if "颜色" not in str(entry.get("spec_name") or "")), None)
            color = next((str(entry.get("spec_value_name") or "") for entry in specs if "颜色" in str(entry.get("spec_name") or "")), "")
            rows.append({
                "brand": brand,
                "item_no": item_no,
                "title": str(product.get("item_name") or ""),
                "category": category,
                "source_size": str((preferred or fallback or {}).get("spec_value_name") or ""),
                "source_color": color,
                "source_price_cents": first_positive(sku.get("activity_price"), sku.get("act_price"), sku.get("price"), product_price),
                "source_market_price_cents": first_positive(sku.get("market_price"), product.get("market_price")),
                "source_stock": stock,
                "region_name": product["_region_name"],
                "source_name": product["_source_name"],
                "source_sku_id": str(sku.get("item_id") or ""),
            })
    return rows


def merge_source_skus(rows):
    merged = {}
    for row in rows:
        key = (row["brand"], row["item_no"], normalize_spec(row["source_size"]), normalize_spec(row["source_color"]))
        if key not in merged:
            merged[key] = {**row, "regions": [row["region_name"]], "sources": [row["source_name"]]}
            continue
        current = merged[key]
        current["source_stock"] += row["source_stock"]
        if row["region_name"] not in current["regions"]:
            current["regions"].append(row["region_name"])
        if row["source_name"] not in current["sources"]:
            current["sources"].append(row["source_name"])
        if row["source_price_cents"] < current["source_price_cents"]:
            for field in ("source_price_cents", "region_name", "source_name", "source_sku_id"):
                current[field] = row[field]
    return list(merged.values())


def atomic_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def compare(source, result):
    if result.get("status") == "error":
        message = result.get("error") or ""
        status = "得物不可出价" if "出价权限" in message else "得物查询失败"
        return {**source, "match_status": status, "error": message, "profitable": False}
    if result.get("status") != "ok":
        return {**source, "match_status": "得物无精确商品", "profitable": False}
    sku = next((entry for entry in result.get("skus", []) if normalize_spec(entry.get("dewu_size")) == normalize_spec(source["source_size"])), None)
    common = {"spu_id": result.get("spu_id"), "dewu_title": result.get("dewu_title"), "total_monthly_sales": result.get("total_monthly_sales", 0)}
    if not sku:
        return {**source, **common, "match_status": "得物尺码未匹配", "profitable": False}
    profit = sku["estimated_net_cents"] - source["source_price_cents"]
    status = "可卖且盈利" if profit > 0 else "可卖但不盈利"
    if not sku["platform_min_price_cents"]:
        status = "得物该尺码缺货"
    elif not sku["bidding_auth"]:
        status = "得物不可出价"
    return {
        **source, **common, **sku,
        "gross_profit_cents": profit,
        "profit_rate": profit / source["source_price_cents"] if source["source_price_cents"] else None,
        "match_status": status,
        "profitable": status == "可卖且盈利",
    }


HEADERS = ["品牌", "货号", "商品名称", "品类", "货源尺码", "颜色", "采购价", "货源库存", "地区", "来源门店", "得物SKU ID", "得物尺码", "平台最低价", "预计到手", "近30天销量", "单件毛利", "成本利润率", "判断/建议", "得物标签"]
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "DCFCE7"
RED = "FEE2E2"
AMBER = "FEF3C7"


def excel_text(value, fallback="未知"):
    text = str(value or "").strip()
    if text.upper().startswith('"ERROR:') or text.upper().startswith("ERROR:"):
        return fallback
    return text


def excel_row(row, profitable_sheet=False):
    advice = "优先考虑" if row.get("profitable") and row.get("monthly_sales", 0) >= 10 else "低销量谨慎" if row.get("profitable") else row.get("match_status")
    return [
        row["brand"], row["item_no"], excel_text(row["title"], "商品名称未知"), row["category"],
        excel_text(row["source_size"]), excel_text(row["source_color"]),
        row["source_price_cents"] / 100, row["source_stock"], "、".join(row.get("regions", [])), "、".join(row.get("sources", [])),
        row.get("dewu_sku_id", ""), row.get("dewu_size", ""), (row.get("platform_min_price_cents") or 0) / 100 or None,
        (row.get("estimated_net_cents") or 0) / 100 or None, row.get("monthly_sales"), None, None,
        advice if profitable_sheet else row.get("match_status"), "、".join(row.get("dewu_tags", [])),
    ]


def add_data_sheet(workbook, name, title, rows, profitable_sheet=False):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = title
    sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="172B4D")
    sheet["A3"] = f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  毛利=得物预计到手-小程序采购价"
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="6B7280")
    header_row = 6
    for column, value in enumerate(HEADERS, 1):
        cell = sheet.cell(header_row, column, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for offset, source in enumerate(rows, header_row + 1):
        for column, value in enumerate(excel_row(source, profitable_sheet), 1):
            sheet.cell(offset, column, value)
        sheet.cell(offset, 16, f'=IF(OR(G{offset}="",N{offset}=""),"",N{offset}-G{offset})')
        sheet.cell(offset, 17, f'=IF(OR(G{offset}="",P{offset}=""),"",P{offset}/G{offset})')
    end_row = max(header_row, header_row + len(rows))
    table = Table(displayName=f"Table_{name.replace(' ', '_').replace('-', '_')}", ref=f"A{header_row}:S{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    sheet.add_table(table)
    sheet.freeze_panes = "C7"
    widths = [15, 16, 42, 22, 11, 18, 13, 11, 12, 24, 16, 11, 14, 14, 13, 13, 13, 16, 18]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=end_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="center")
    for column in (7, 13, 14, 16):
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=header_row + 1, max_row=end_row):
            for item in cell:
                item.number_format = "#,##0.00"
    for cell in sheet.iter_cols(min_col=17, max_col=17, min_row=header_row + 1, max_row=end_row):
        for item in cell:
            item.number_format = "0.0%"
    if rows:
        sheet.conditional_formatting.add(f"P{header_row + 1}:P{end_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
        sheet.conditional_formatting.add(f"P{header_row + 1}:P{end_row}", CellIsRule(operator="lessThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
        sheet.conditional_formatting.add(f"R{header_row + 1}:R{end_row}", FormulaRule(formula=[f'ISNUMBER(SEARCH("优先",R{header_row + 1}))'], fill=PatternFill("solid", fgColor=GREEN)))
        sheet.conditional_formatting.add(f"R{header_row + 1}:R{end_row}", FormulaRule(formula=[f'ISNUMBER(SEARCH("未匹配",R{header_row + 1}))'], fill=PatternFill("solid", fgColor=AMBER)))
    return sheet


def build_workbook(output: Path, rows, profile_counts):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("汇总")
    summary.sheet_view.showGridLines = False
    summary["A2"] = "ADIDAS 与 ANTA KIDS 得物盈利分析"
    summary["A2"].font = Font(name="Arial", size=16, bold=True, color="172B4D")
    summary.append([])
    summary.append(["品牌", "实时货号", "货源SKU", "匹配SKU", "盈利SKU", "盈利货号", "优先考虑SKU"])
    for cell in summary[4]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for brand in TARGET_PROFILES:
        brand_rows = [row for row in rows if row["brand"] == brand]
        summary.append([
            brand,
            len({row["item_no"] for row in brand_rows}),
            len(brand_rows),
            sum(bool(row.get("dewu_sku_id")) for row in brand_rows),
            sum(bool(row.get("profitable")) for row in brand_rows),
            len({row["item_no"] for row in brand_rows if row.get("profitable")}),
            sum(bool(row.get("profitable")) and row.get("monthly_sales", 0) >= 10 for row in brand_rows),
        ])
    summary.append([])
    summary.append(["门店实时结果", "商品数"])
    for profile_name, members in profile_counts.items():
        for member in members:
            summary.append([f"{profile_name} / {member['regionName']} / {member['sourceName']}", member["productCount"]])
    for width, column in zip((50, 16, 16, 16, 16, 16, 16), "ABCDEFG"):
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A4"

    all_unmatched = []
    for brand, prefix in (("ADIDAS-合并", "ADIDAS"), ("ANTA KIDS - 聚合", "ANTA KIDS")):
        brand_rows = [row for row in rows if row["brand"] == brand]
        profitable = sorted((row for row in brand_rows if row.get("profitable")), key=lambda row: (row.get("monthly_sales", 0) < 10, -row.get("monthly_sales", 0), -row.get("gross_profit_cents", 0)))
        add_data_sheet(workbook, f"{prefix}盈利", f"{prefix} 盈利机会", profitable, True)
        add_data_sheet(workbook, f"{prefix}全部", f"{prefix} 全部货源与得物对比", brand_rows)
        all_unmatched.extend(row for row in brand_rows if row.get("match_status") in ("得物无精确商品", "得物尺码未匹配"))
    add_data_sheet(workbook, "未匹配", "未匹配商品与尺码", all_unmatched)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="outputs/ADIDAS_ANTA_KIDS_得物盈利分析.xlsx")
    parser.add_argument("--cache", default="runtime/dewu-adidas-anta-profit-cache.json")
    parser.add_argument("--interval", type=float, default=None, help="兼容参数：固定请求间隔")
    parser.add_argument("--min-interval", type=float, default=2.0)
    parser.add_argument("--max-interval", type=float, default=4.0)
    parser.add_argument("--source-cache", default="runtime/dewu-adidas-anta-source-cache.json")
    parser.add_argument("--refresh-source", action="store_true")
    parser.add_argument("--browser-session", help="在指定 agent-browser 会话中执行得物 API")
    parser.add_argument("--verify-timeout", type=int, default=600, help="等待手动完成得物校验的秒数")
    parser.add_argument("--browser-batch-size", type=int, default=10, help="浏览器模式每批货号数")
    args = parser.parse_args()
    output = Path(args.output).resolve()
    cache_path = Path(args.cache).resolve()
    source_cache_path = Path(args.source_cache).resolve()
    cache = json.loads(cache_path.read_text(encoding="utf-8")) if cache_path.exists() else {}

    if source_cache_path.exists() and not args.refresh_source:
        source_cache = json.loads(source_cache_path.read_text(encoding="utf-8"))
        merged = source_cache["merged"]
        profile_counts = source_cache["profile_counts"]
        print(f"复用实时货源缓存：{source_cache_path}", flush=True)
    else:
        print("读取聚合品牌实时货源...", flush=True)
        profiles = load_profiles()
        all_source = []
        profile_counts = {}
        for profile in profiles:
            print(profile["name"], flush=True)
            products, counts = load_profile_products(profile)
            profile_counts[profile["name"]] = counts
            all_source.extend(source_skus(profile["name"], products))
        merged = merge_source_skus(all_source)
        atomic_json(source_cache_path, {"saved_at": datetime.now().isoformat(), "merged": merged, "profile_counts": profile_counts})
    item_nos = sorted({row["item_no"] for row in merged})
    print(f"合并后 {len(item_nos)} 个货号、{len(merged)} 个有库存 SKU", flush=True)

    fixed_interval = args.interval
    min_interval = fixed_interval if fixed_interval is not None else args.min_interval
    max_interval = fixed_interval if fixed_interval is not None else args.max_interval
    if args.browser_session:
        print(f"使用真实浏览器会话：{args.browser_session}", flush=True)
        client = BrowserDewuClient(min_interval, max_interval, args.browser_session, args.verify_timeout)
    else:
        client = DewuClient(min_interval, max_interval)
    try:
        if isinstance(client, BrowserDewuClient):
            pending = [
                item_no
                for item_no in item_nos
                if cache.get(item_no, {}).get("status") not in ("ok", "product_not_found", "error")
            ]
            batch_size = max(1, args.browser_batch_size)
            print(f"待查询 {len(pending)} 个货号，每批 {batch_size} 个", flush=True)
            for offset in range(0, len(pending), batch_size):
                batch = pending[offset:offset + batch_size]
                for attempt in range(1, 4):
                    try:
                        cache.update(client.query_items(batch))
                        break
                    except ValidationRequired as error:
                        raise ValidationRequired(f"批次 {batch[0]} - {batch[-1]}：{error}") from error
                    except Exception as error:
                        if attempt == 3:
                            atomic_json(cache_path, cache)
                            raise RuntimeError(
                                f"批次 {batch[0]} - {batch[-1]} 连续失败，未写入错误缓存：{error}"
                            ) from error
                        print(f"批次请求失败，第 {attempt}/3 次重试：{error}", file=sys.stderr, flush=True)
                        time.sleep(2)
                atomic_json(cache_path, cache)
                done = min(offset + len(batch), len(pending))
                print(f"本次查询进度 {done}/{len(pending)}；总缓存 {len(cache)}/{len(item_nos)}", flush=True)
        else:
            for index, item_no in enumerate(item_nos, 1):
                if cache.get(item_no, {}).get("status") not in ("ok", "product_not_found", "error"):
                    try:
                        cache[item_no] = client.query_item(item_no)
                    except ValidationRequired as error:
                        raise ValidationRequired(f"货号 {item_no}：{error}") from error
                    except Exception as error:
                        cache[item_no] = {"status": "error", "error": str(error)}
                    atomic_json(cache_path, cache)
                if index % 10 == 0 or index == len(item_nos):
                    print(f"得物查询进度 {index}/{len(item_nos)}", flush=True)
    except ValidationRequired as error:
        atomic_json(cache_path, cache)
        print(str(error), file=sys.stderr, flush=True)
        print(f"已保存缓存：{cache_path}", file=sys.stderr, flush=True)
        return 85

    compared = [compare(row, cache.get(row["item_no"], {})) for row in merged]
    build_workbook(output, compared, profile_counts)
    counts = Counter(row["match_status"] for row in compared)
    print(json.dumps({"output": str(output), "rows": len(compared), "status": counts}, ensure_ascii=False), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
